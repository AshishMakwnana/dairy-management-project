from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "projects assests" / "documents" / "DUDHO_Master_UI_UX_Development_Blueprint_v3.docx"
OUTPUT = ROOT / "projects assests" / "documents" / "DUDHO_Master_UI_UX_Development_Blueprint_v4.docx"


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_borders(cell, color: str = "D9D9D9") -> None:
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), color)


def add_bullet(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(text, style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(3)


document = Document(SOURCE)

document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
heading = document.add_paragraph("Phase 0 Implementation Record", style="Heading 1")
for run in heading.runs:
    run.font.color.rgb = RGBColor(0, 0, 0)

intro = document.add_paragraph()
intro.add_run("Status: ").bold = True
intro.add_run("In progress")
intro.add_run("    Started: ").bold = True
intro.add_run("9 October 2026")

document.add_heading("Purpose and outcome", level=2)
document.add_paragraph(
    "Phase 0 establishes the development environments, application boundaries, versioned API foundation, "
    "verification workflow and shared documentation required before authentication and dairy onboarding begin. "
    "The Laravel and React Native applications now exist, the backend infrastructure is operational, and the "
    "mobile foundation can verify availability of the versioned Laravel API."
)

document.add_heading("Confirmed architecture", level=2)
for item in (
    "The Expo React Native application runs on Windows outside Docker.",
    "Laravel, Nginx, MySQL, Redis, Mailpit, the queue worker and the scheduler run in Docker Engine inside WSL2 Ubuntu.",
    "The local API uses /api/v1 and is exposed through Nginx on port 8082. Port 8080 remains assigned to an unrelated local project.",
    "Laravel remains authoritative for authentication, authorization, pricing and final financial values.",
    "Initial quality values are entered manually. Physical analyzer, Bluetooth, USB and vendor SDK integration are deferred.",
    "Mobile owner workflows are the first delivery priority. Laravel web administration is required later and will reuse the same services and policies.",
):
    add_bullet(document, item)

document.add_heading("Implemented foundation", level=2)
table = document.add_table(rows=1, cols=3)
table.style = "Table Grid"
table.autofit = True
headers = ("Component", "Implementation", "Verification")
for index, value in enumerate(headers):
    cell = table.rows[0].cells[index]
    cell.text = value
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_shading(cell, "D9EAF7")
    set_cell_borders(cell)
    for run in cell.paragraphs[0].runs:
        run.bold = True
        run.font.color.rgb = RGBColor(0, 0, 0)

rows = (
    ("Backend", "Laravel 13 with PHP 8.3 target", "Generated and feature tests passed"),
    ("API", "GET /api/v1/health with standard response envelope", "HTTP request through Nginx passed"),
    ("Database", "MySQL 8.4 in Docker", "Initial migrations passed"),
    ("Infrastructure", "Nginx, Redis, Mailpit, queue worker and scheduler", "Docker services built and started"),
    ("Mobile", "Expo SDK 57, React Native 0.86 and TypeScript", "TypeScript compilation passed"),
    ("Mobile integration", "Startup connectivity status using EXPO_PUBLIC_API_URL override", "Expo web bundle passed"),
    ("Runtime", "Node 22.13.0", "React Native engine requirement satisfied"),
    ("Agent collaboration", "Root Git repository, agent instructions, context, task board and handoff log", "Workspace context created"),
)

for row_index, row_values in enumerate(rows, start=1):
    cells = table.add_row().cells
    for column_index, value in enumerate(row_values):
        cell = cells[column_index]
        cell.text = value
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_borders(cell)
        if row_index % 2 == 0:
            set_cell_shading(cell, "F8FAFC")

document.add_heading("Development endpoints", level=2)
for item in (
    "Laravel API from Windows or browser: http://localhost:8082/api/v1",
    "Laravel API from an Android emulator on a supported host: http://10.0.2.2:8082/api/v1",
    "Mailpit inbox: http://localhost:8025",
    "MySQL host connection: 127.0.0.1:3307",
):
    add_bullet(document, item)

document.add_heading("Constraints and open decisions", level=2)
for item in (
    "The development computer uses Windows on ARM64. Google does not currently support Android Studio and Android Emulator execution on this host. Android validation requires a physical Android phone or Android Device Streaming.",
    "The final brand direction remains unresolved because the blueprint selects a blue D-shaped mark while a production-named asset uses a cartoon farmer and cow identity.",
    "SMS and WhatsApp OTP providers have not been selected.",
    "The complete Phase 1 API contract and authorization matrix must be approved before authentication implementation begins.",
    "Detailed pricing formulas, precision and rounding rules remain Phase 2 product decisions.",
    "The generated Expo dependency tree reports transitive advisories. Forced npm remediation is prohibited because it proposes incompatible framework downgrades; upstream-compatible releases will be monitored.",
):
    add_bullet(document, item)

document.add_heading("Phase 0 exit gate", level=2)
document.add_paragraph(
    "Phase 0 will be complete after the mobile application is validated on a supported Android target, the mobile-to-API connection is confirmed on that target, the final brand direction is approved, and the Phase 1 API and authorization contracts are documented."
)

document.add_heading("Next implementation package", level=2)
document.add_paragraph(
    "The next package is Phase 1 Authentication and Dairy Onboarding. Work begins with the exact API contract, "
    "Laravel Sanctum, the Phase 1 migrations, OTP provider abstraction, authorization policies, shared mobile "
    "theme and localization foundations, and the login and OTP vertical slice."
)

for style_name in ("Heading 1", "Heading 2", "Heading 3"):
    style = document.styles[style_name]
    style.font.color.rgb = RGBColor(0, 0, 0)

document.save(OUTPUT)
print(OUTPUT)
